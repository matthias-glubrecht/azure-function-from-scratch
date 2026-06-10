"""
build-slides.py
===============
Reproduzierbarer Generator für slides.pptx — Vortrag
"Azure Function von Grund auf (TypeScript) — am Beispiel ManagePermissions".

Begleitdeck zur ausführlichen anleitung.md im selben Ordner. Stil, Theme und
Hilfsfunktionen sind bewusst identisch zum Haupt-Workshop-Deck
(../praesentation/build-slides.py), damit beide Decks zusammenpassen.

Voraussetzungen:
    pip install python-pptx

Aufruf:
    python build-slides.py
        -> erzeugt slides.pptx im selben Verzeichnis
    python build-slides.py --output slides.new.pptx
        -> alternative Datei (nützlich, wenn slides.pptx in PowerPoint geöffnet/gesperrt ist)
"""

import argparse
import re
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

# ---------------------------------------------------------------------------
# Slide-Definitionen
# ---------------------------------------------------------------------------
# Pro Folie ein Dict mit:
#   title   — Überschrift (bzw. Haupttitel bei Folie 1)
#   bullets — Body-Zeilen (oder Subtitle-Zeilen bei Folie 1; bei Tabellen: Caption)
#   notes   — Speaker Notes für den Trainer (Stichpunkte)
#   table   — optional {"columns": [...], "rows": [[...], ...]}
#
# Konvention: Zeilen, die mit 3 Leerzeichen ODER "→" beginnen, werden als
# Sub-Bullet (eingerückt, gedämpft) gerendert.
SLIDES = [
    # 1. Title-Slide
    {
        "title": "Azure Function von Grund auf",
        "bullets": [
            "Vom leeren Ordner zur laufenden Cloud-Function — von Hand, ohne KI",
            "TypeScript / Node.js · PnPjs · Managed Identity",
            "Durchgehendes Beispiel: die ManagePermissions-Lösung",
            "So konkret wie ein SPFx-Web-Part, das man selbst aufsetzt",
        ],
        "notes": [
            "Rahmung: Heute bauen wir eine Azure Function von Grund auf selbst — kein Generator-im-Hintergrund, kein KI-Autopilot. Genau das, was ihr bei einem SPFx-Web-Part längst könnt, machen wir jetzt für eine Function.",
            "Beispiel ist die ManagePermissions-Lösung — die gibt es produktiv in C#; wir bauen sie hier in TypeScript nach, um die Schritte zu zeigen.",
            "Versprechen an die Teilnehmer: Am Ende könnt ihr vor einem leeren Ordner sitzen und wisst, welches Werkzeug ihr braucht und welcher Schritt als Nächstes kommt.",
            "Begleitend gibt es anleitung.md zum Mitmachen — Befehl für Befehl, mit komplettem Code.",
        ],
    },

    # 2. Worum es geht
    {
        "title": "Worum es geht",
        "bullets": [
            "Frage: Wie baue ich eine Azure Function, wenn ich es NICHT die KI machen lasse?",
            "Bei SPFx weißt du das längst — hier dieselbe Sicherheit für Functions",
            "Erst Voraussetzungen (Werkzeuge), dann Schritte nacheinander",
            "Keine Vorkenntnisse über Azure Functions nötig",
            "Ziel ist das übertragbare Muster, nicht das einzelne Sample",
        ],
        "notes": [
            "Das ehrliche Motiv: Man hat zuletzt viel automatisch erledigen lassen. Wer einen Weg einmal selbst gegangen ist, versteht ihn — und kann ihn debuggen, erklären, anpassen.",
            "Analogie SPFx: yeoman-Generator, gulp serve, package-solution, App-Katalog — diesen Ablauf habt ihr im Kopf. Wir bauen genau dieses mentale Modell für Functions auf.",
            "Aufbau des Vortrags ansagen: 1) Denkmodell, 2) Was wir bauen, 3) Voraussetzungen, 4) zwölf Schritte.",
            "Wichtig: Es geht nicht darum, ManagePermissions auswendig zu lernen, sondern den Weg 'leerer Ordner -> Cloud-Function' zu beherrschen.",
        ],
    },

    # 3. Denkmodell SPFx <-> Functions (Tabelle)
    {
        "title": "Denkmodell: SPFx-Web-Part ↔ Azure Function",
        "bullets": [
            "Du kennst den Ablauf bereits — nur die Werkzeuge heißen anders:",
        ],
        "table": {
            "columns": ["Bei SPFx", "Bei Azure Functions"],
            "rows": [
                ["yo @microsoft/sharepoint (scaffolden)", "func init (scaffolden)"],
                ["Web-Part-Komponente hinzufügen", "func new (Function/Trigger hinzufügen)"],
                ["gulp serve (lokaler Workbench)", "func start (Host auf localhost:7071)"],
                ["package-solution.json", "host.json (App-weite Konfiguration)"],
                [".sppkg in den App-Katalog", "func azure functionapp publish"],
                ["webApiPermissionRequests", "Entra-App + Sites.Selected-Grant"],
                ["Properties im Property-Pane", "App Settings (Umgebungsvariablen)"],
                ["SPHttpClient / AadHttpClient", "PnPjs + DefaultAzureCredential"],
            ],
        },
        "notes": [
            "Diese Tabelle ist der rote Faden des ganzen Vortrags. Jede Zeile ist ein 'Aha, das kenne ich ja schon'.",
            "Kernunterschied unten betonen: SPFx läuft im Browser MIT der Identität des Benutzers. Die Function läuft serverseitig und braucht eine EIGENE Identität (Managed Identity), um auf SharePoint zuzugreifen.",
            "Wer die linke Spalte beherrscht, muss für die rechte nur die neuen Befehlsnamen lernen — die Konzepte sind dieselben.",
        ],
    },

    # 4. Was wir bauen
    {
        "title": "Was wir bauen — ManagePermissions",
        "bullets": [
            "Ein HTTP-Endpunkt: POST /api/ManagePermissions",
            "Setzt/entfernt Berechtigung auf EINEM SharePoint-Listenelement",
            "Aufgerufen aus einem SPFx-Web-Part",
            "action = grant | reset, dazu webUrl, listId, itemId, …",
            "Antwort: 200 { ok, message } oder nicht-2xx { error }",
        ],
        "notes": [
            "Use-Case malen: Ein Sachbearbeiter klickt im Web-Part 'Diesem Kollegen Zugriff auf dieses eine Dokument geben' — und die Function erledigt genau das, kontrolliert.",
            "Bewusst klein gehalten: EIN Endpunkt, ZWEI Aktionen (grant/reset). Klein genug zum Durchspielen, aber es braucht alle relevanten Bausteine.",
            "grant = Vererbung trennen + Rolle setzen. reset = Vererbung wiederherstellen. Mehr macht die Function nicht.",
        ],
    },

    # 5. Zwei Vertrauensgrenzen
    {
        "title": "Zwei Vertrauensgrenzen, zwei Identitäten",
        "bullets": [
            "1) Aufrufer → Function = delegated",
            "   Benutzer-Token (access_as_user), JWT im Code geprüft + Gruppen-Check",
            "2) Function → SharePoint = app-only",
            "   Managed Identity, Sites.Selected + FullControl (nur auf der Ziel-Site)",
            "Kontrollierte Rechte-Erhöhung: der Aufrufer braucht selbst KEINE Site-Admin-Rechte",
        ],
        "notes": [
            "Das ist die wichtigste konzeptionelle Folie. Zwei Grenzen, zwei völlig verschiedene Identitäten — das verwechseln Einsteiger am häufigsten.",
            "Grenze 1 (delegated): 'Wer ruft an, und darf der das?' Das Token gehört dem angemeldeten Benutzer. Wir prüfen Signatur, Issuer, Audience, Gültigkeit, Scope und Gruppenmitgliedschaft.",
            "Grenze 2 (app-only): 'Mit welchem Recht fasse ich SharePoint an?' Hier zählt die Identität der Function App selbst — die Managed Identity. Least privilege: nur Sites.Selected, und nur auf der einen Ziel-Site.",
            "Pointe: Der Aufrufer kann Rechte vergeben, die er selbst gar nicht hat. Genau deshalb ist die Autorisierung in Grenze 1 (Gruppe!) so wichtig.",
        ],
    },

    # 6. Voraussetzungen — Werkzeuge (Tabelle)
    {
        "title": "Voraussetzungen — Werkzeuge",
        "bullets": [
            "Einmalig installieren (winget / npm):",
        ],
        "table": {
            "columns": ["Werkzeug", "Zweck"],
            "rows": [
                ["Node.js 22 LTS", "Laufzeit + npm; die App läuft später auf Node 22"],
                ["Azure Functions Core Tools 4.x", "func: init / new / start / publish"],
                ["Azure CLI 2.60+", "Ressourcen anlegen; az login (lokal = Identität)"],
                ["Git", "Versionsverwaltung"],
                ["Azurite", "Lokaler Storage-Emulator für func start"],
                ["VS Code + Azure Functions Ext.", "Optional, bequemes Debuggen/Deployen"],
                ["PnP.PowerShell + Microsoft.Graph", "Nur fürs MI-Berechtigen (Schritt 11)"],
            ],
        },
        "notes": [
            "winget install OpenJS.NodeJS.LTS / Microsoft.AzureCLI / Microsoft.Azure.FunctionsCoreTools; npm install -g azurite.",
            "Kontrolle vorführen: node --version (v22), func --version (4.x), az --version.",
            "Azurite kurz erklären: Auch lokal will der Functions-Host einen Storage-Account für interne Buchführung. Azurite emuliert den, damit func start nicht meckert.",
            "Die PowerShell-Module braucht man erst ganz am Ende (MI berechtigen) — nicht abschrecken lassen, der Code selbst braucht kein PowerShell.",
        ],
    },

    # 7. Voraussetzungen — Konten & Rollen
    {
        "title": "Voraussetzungen — Konten & Rollen",
        "bullets": [
            "Azure-Subscription mit Contributor (Ressourcen + Function App anlegen)",
            "Entra-Tenant-Admin — einmalig für den Admin-Consent von Sites.Selected",
            "SharePoint-Administrator — Ziel-Site für die MI freigeben",
            "   und ggf. die API-Berechtigung des SPFx-Web-Parts genehmigen",
        ],
        "notes": [
            "Ehrlich sein: Ohne diese drei Rollen kommt man an zwei Stellen nicht weiter (Consent + Site-Grant). Das früh klären, sonst hängt man im Workshop.",
            "Contributor reicht für die Azure-Seite — man braucht NICHT Owner.",
            "Der Entra-Admin-Consent für Sites.Selected ist ein EINMALIGER Schritt pro Tenant.",
            "Der SharePoint-Admin gibt der Managed Identity später FullControl auf genau EINER Site — nicht tenantweit. Das ist der least-privilege-Clou.",
        ],
    },

    # 8. Schritt 1 — func init
    {
        "title": "Schritt 1 — func init",
        "bullets": [
            "mkdir managepermissions-ts && cd … && git init",
            "func init --worker-runtime node --language typescript",
            "   --worker-runtime = Sprachwahl: node, python, powershell, dotnet-isolated, java, custom",
            "Das ist das func-Gegenstück zu yo @microsoft/sharepoint",
            "Aus dem leeren Ordner wird ein vollständiges Function-Projekt",
        ],
        "notes": [
            "Hier passiert der titelgebende Moment: leerer Ordner -> Projekt. Genau die Frage des Teilnehmers.",
            "func init legt das Gerüst an: host.json, local.settings.json, package.json, tsconfig.json, .funcignore, .gitignore, src/.",
            "--worker-runtime (nur informativ, falls gefragt): legt die Sprache fest, in der ihr Functions schreibt. node = JS/TS (mit --language javascript/typescript), python, powershell, dotnet-isolated = C#/F# (aktueller .NET-Standard), dotnet = C# in-process (älter, läuft aus), java, custom = beliebige Sprache über HTTP (z.B. Go/Rust). Die produktive ManagePermissions nutzt dotnet-isolated; wir nehmen node.",
            "Betonen: Bis hierher ist NICHTS Azure-spezifisch passiert außer Dateien anlegen. Kein Cloud-Login nötig.",
            "Nächste Folie zeigt, was diese Dateien bedeuten.",
        ],
    },

    # 9. Die generierten Dateien (Tabelle)
    {
        "title": "Die generierten Dateien",
        "bullets": [
            "Was func init anlegt — und was es bedeutet:",
        ],
        "table": {
            "columns": ["Datei", "Bedeutung"],
            "rows": [
                ["host.json", "App-weite Konfiguration (alle Functions)"],
                ["local.settings.json", "Lokale Settings/Secrets; wird NICHT deployt"],
                ["package.json", "npm-Pakete + Scripts; main = wo Functions registriert sind"],
                ["tsconfig.json", "TypeScript-Compiler: src/**/*.ts → dist/"],
                [".funcignore", "Was NICHT mit deployt wird (analog .gitignore)"],
                [".gitignore", "Schließt local.settings.json, node_modules, dist aus"],
            ],
        },
        "notes": [
            "Das main-Feld in package.json ist der wichtigste Aha-Punkt: 'dist/src/functions/*.js' — der Host lädt diese Dateien beim Start und führt deren Registrierungen aus.",
            "local.settings.json hervorheben: Hier stehen lokale Geheimnisse, und genau deshalb ist die Datei in .gitignore und im .funcignore — sie landet weder im Git noch in Azure.",
            "tsconfig: src wird nach dist kompiliert. Deshalb zeigt main auf dist/, nicht auf src/.",
        ],
    },

    # 10. Schritt 2 — func new + lokal starten
    {
        "title": "Schritt 2 — func new + lokal starten",
        "bullets": [
            'func new --name ManagePermissions --template "HTTP trigger" --authlevel anonymous',
            "   --authlevel: anonymous | function | admin",
            "→ erzeugt src/functions/ManagePermissions.ts (Hello world)",
            "npm install && npm start  →  Host auf http://localhost:7071",
            "Invoke-RestMethod …/api/ManagePermissions?name=Workshop → Hello, Workshop!",
            "Meilenstein: eine Function läuft lokal — ab jetzt nur noch Logik ersetzen",
        ],
        "notes": [
            "func new = 'Komponente hinzufügen'. authLevel anonymous, WEIL wir die Identität selbst im Code prüfen (Token+Gruppe) — der Function-API-Key würde das nur verdoppeln.",
            "authLevel steuert, ob ein von Azure generierter Function-Schlüssel mitgeschickt werden muss. Drei Werte: anonymous = kein Schlüssel (wenn ihr die Auth selbst macht oder ein Türsteher wie API Management davorsteht); function = Function-/Host-Schlüssel nötig (Standard, für interne/Service-zu-Service-Aufrufe); admin = Master-Schlüssel (sehr weitreichend, meiden). WICHTIG: anonymous heißt NICHT ungeschützt — nur 'kein Azure-Schlüssel'; unser JWT-Check ist der eigentliche Türsteher.",
            "npm start läuft prestart -> clean+build (tsc) -> func start. Den ersten erfolgreichen Aufruf gemeinsam feiern: Das ist der 'es lebt'-Moment.",
            "Pädagogisch wichtig: Erst das triviale Hello-World zum Laufen bringen, DANN Komplexität dazu. So weiß man bei jedem späteren Fehler, dass das Gerüst stand.",
            "Falls Storage-Warnung: azurite starten oder AzureWebJobsStorage=UseDevelopmentStorage=true.",
        ],
    },

    # 11. v4-Modell
    {
        "title": "Das v4-Programmiermodell — Functions im Code",
        "bullets": [
            "'Programmiermodell' = WIE man Functions im Code schreibt",
            "v4 = aktuelle Node-Generation (npm-Paket @azure/functions v4)",
            "Function 'registrieren' = dem Host bekannt machen:",
            "   app.http('ManagePermissions', { methods, authLevel, handler })",
            "   → Name · worauf sie reagiert · welcher Code läuft",
            "Beim Start lädt der Host die Dateien aus main und führt die Registrierungen aus",
            "Pendant zu C#: [Function(\"…\")] / [HttpTrigger] im Code — gleiche Idee",
        ],
        "notes": [
            "Begriff entschärfen: 'Programmiermodell' klingt groß, meint aber nur die ART, WIE man im Code eine Function beschreibt — welche Datei-Struktur, welche Befehle. Für Node ist v4 die aktuelle, einsteigerfreundliche Generation; sie hängt am npm-Paket @azure/functions Version 4.",
            "'Registrieren' in einfachen Worten: app.http(...) meldet die Function beim Host AN. Du sagst ihm drei Dinge — wie sie heißt ('ManagePermissions'), worauf sie reagiert (HTTP GET/POST) und welche Funktion laufen soll (handler). Mehr ist es nicht.",
            "Ablauf beim Start: Der Host liest die Dateien aus dem main-Feld der package.json, führt die app.http(...)-Aufrufe darin aus und weiß danach, welche Endpunkte es gibt und wohin Anfragen müssen.",
            "Deshalb MUSS app.http(...) auf globaler Ebene stehen (nicht in einer anderen Funktion) — sonst wird es beim Laden nicht ausgeführt und der Host kennt die Function nicht.",
            "Nur FALLS online ältere Beispiele auftauchen: Früher lag neben jeder Function eine separate function.json. Im v4-Modell braucht man sie nicht mehr — alles steht im Code. Nicht von sich aus mit function.json anfangen, wenn im Raum niemand sie kennt.",
            "Schöne Parallele zu C# isolated: Dort macht das Attribut [Function(\"…\")] / [HttpTrigger] dasselbe wie app.http(...). Gleiche Designidee, zwei Sprachen.",
        ],
    },

    # 12. Schritt 3 — Abhängigkeiten (Tabelle)
    {
        "title": "Schritt 3 — Abhängigkeiten",
        "bullets": [
            "npm install @azure/identity @pnp/sp @pnp/nodejs @pnp/azidjsclient jose",
        ],
        "table": {
            "columns": ["Paket", "Wofür"],
            "rows": [
                ["@azure/identity", "DefaultAzureCredential (MI in Azure, az-Login lokal)"],
                ["@pnp/sp", "PnPjs-Kern: Webs, Listen, Items, Security"],
                ["@pnp/nodejs", "Node-Bausteine: SPDefault (fetch-Client)"],
                ["@pnp/azidjsclient", "Brücke @azure/identity → PnPjs (AzureIdentity)"],
                ["jose", "JWT-Validierung (Signaturprüfung gegen Entra-Schlüssel)"],
            ],
        },
        "notes": [
            "Diese fünf Pakete sind das ganze Fundament. Jedes hat genau eine Aufgabe — keine Magie.",
            "@pnp/azidjsclient ist der Klebstoff: Es macht aus einer @azure/identity-Credential ein PnPjs-Auth-Behavior. Damit muss man KEINEN Token-Code von Hand schreiben.",
            "jose statt selbstgebauter Krypto: Token-Validierung will man NICHT von Hand machen. jose prüft Signatur gegen die öffentlichen Schlüssel von Entra und cacht sie automatisch.",
            "Dateistruktur ansagen: functions/ (Handler), auth/ (Token), services/ (PnPjs), models/ (Typen) — wie in der C#-Lösung.",
        ],
    },

    # 13. Schritt 4 — Eingabe-Vertrag
    {
        "title": "Schritt 4 — Der Eingabe-Vertrag",
        "bullets": [
            "models/managePermissionsRequest.ts — ein TypeScript-Interface",
            "action, webUrl, listId, itemId",
            "userPrincipalName, permissionLevel (nur grant)",
            "copyExistingPermissions (nur grant, Default true)",
            "Der Vertrag dokumentiert, was das Web-Part schicken muss",
        ],
        "notes": [
            "Ein Interface ist hier reine Dokumentation + Typsicherheit — zur Laufzeit prüfen wir die Werte trotzdem selbst (kommt im Service).",
            "permissionLevel ist eine Allowlist: Read, Contribute, Edit, Design, FullControl. Mehr lassen wir nicht zu.",
            "copyExistingPermissions=false als interessante Option: Element startet exklusiv, nur der neue Benutzer (plus Site-Admins) hat Zugriff. Default true erhält bestehende Zugriffe.",
        ],
    },

    # 14. Schritt 5 — Aufrufer prüfen
    {
        "title": "Schritt 5 — Aufrufer prüfen (JWT im Code)",
        "bullets": [
            "auth/callerAuthorizer.ts mit jose",
            "createRemoteJWKSet(...) lädt Entras öffentliche Schlüssel (gecacht)",
            "jwtVerify: Signatur + Issuer (/v2.0) + Audience + Gültigkeit",
            "Dann scp prüfen (access_as_user) und groups (AllowedGroupId)",
            "Gefälschtes/abgelaufenes/fremdes Token → 401/403",
        ],
        "notes": [
            "Wir machen die Token-Prüfung bewusst SICHTBAR im Code, statt sie hinter Middleware zu verstecken — damit jeder sieht, was 'ein Token validieren' eigentlich heißt.",
            "Vier Dinge prüft die Signatur-/Claim-Validierung: Ist das Token echt (Signatur)? Vom richtigen Aussteller (Issuer v2.0)? Für UNS bestimmt (Audience)? Noch gültig (Lifetime)?",
            "scp = der delegierte Scope, dem der Benutzer zugestimmt hat. groups = unsere Autorisierung: nur Mitglieder einer bestimmten Sicherheitsgruppe dürfen rein.",
            "Frage, die garantiert kommt: groups-Claim ODER roles-Claim? Wir prüfen bewusst groups (Mitgliedschaft in einer Entra-Sicherheitsgruppe, wie die C#-Lösung). Alternative wäre der roles-Claim über App-Rollen: eigene Rollen in der App-Registrierung definieren und Benutzer/Gruppen zuweisen -> im Token erscheint roles. Faustregel: groups, wenn passende Sicherheitsgruppen schon existieren; App-Rollen (roles) sind app-spezifischer, sauberer und vermeiden Group-Overage. Im Workshop bleiben wir bei groups, damit Anleitung und C#-Lösung deckungsgleich sind.",
            "createRemoteJWKSet auf Modulebene anlegen — es cacht die Schlüssel und behandelt Schlüssel-Rotation automatisch. Nicht pro Aufruf neu bauen.",
        ],
    },

    # 15. Zwei Entra-Stolpersteine
    {
        "title": "Zwei Entra-Stolpersteine",
        "bullets": [
            "1) Token-Version: api://-Apps stellen per Default v1.0-Tokens aus",
            "   Issuer passt dann nicht zu /v2.0 → JEDER Aufruf 401",
            "   Fix: requestedAccessTokenVersion = 2 im App-Manifest",
            "2) groups-Claim ist standardmäßig NICHT im Token",
            "   Fix: Token configuration → groups claim → 'Groups assigned to the application'",
        ],
        "notes": [
            "Diese beiden Stolpersteine kosten erfahrungsgemäß die meiste Debugging-Zeit — deshalb eine eigene Folie. Beide äußern sich als 'es geht einfach nicht', obwohl der Code stimmt.",
            "Token v2: Ohne die Umstellung kommt ein v1.0-Issuer (sts.windows.net) statt login.microsoftonline.com/.../v2.0 — und unsere Validierung lehnt korrekt ab. Kontrolle über jwt.ms: Claim ver muss 2.0 sein.",
            "groups-Claim: Standardmäßig steht NICHT im Token, in welchen Gruppen jemand ist. Erst der explizit aktivierte Claim bringt ihn rein. 'Groups assigned to the application' vermeidet außerdem das Overage-Problem.",
            "Lehrsatz: Wenn 'alles richtig' aussieht und trotzdem 401/403 kommt — fast immer einer dieser beiden Schalter in der App-Registrierung.",
        ],
    },

    # 16. Schritt 6 — SharePoint mit PnPjs
    {
        "title": "Schritt 6 — SharePoint mit PnPjs (app-only)",
        "bullets": [
            "services/sharePointPermissionService.ts",
            "const credential = new DefaultAzureCredential()  // einmal, Modulebene",
            "spfi(webUrl).using(SPDefault(), AzureIdentity(credential, [scope], null))",
            "scope = https://<host>/.default  →  SharePoint-Token (nicht Graph)",
            "In Azure = Managed Identity, lokal = az login — KEIN Secret-Code",
        ],
        "notes": [
            "Das ist der Auth-Kern der zweiten Vertrauensgrenze — und erstaunlich wenig Code. DefaultAzureCredential + AzureIdentity erledigen Token-Beschaffung und -Refresh komplett.",
            "DefaultAzureCredential ist clever: dieselbe Codezeile nimmt in Azure die Managed Identity und auf eurem Rechner eure az-login-Anmeldung. Lokal testen ohne Secrets.",
            "Scope .../.default holt ein SharePoint-Token (Audience = der SP-Host). Deshalb braucht die MI nur das SharePoint-Recht Sites.Selected — keine Graph-Rechte.",
            "Best Practice: die Credential EINMAL auf Modulebene anlegen (Token-Cache), nicht pro Aufruf neu.",
        ],
    },

    # 17. Die Berechtigungs-Logik
    {
        "title": "Die Berechtigungs-Logik (grant / reset)",
        "bullets": [
            "roleDefinitions.getByType(roleType) — sprachunabhängig (kein 'Mitwirken'-Problem)",
            "ensureUser(upn) — Benutzer in der Site sicherstellen → userId",
            "breakRoleInheritance(copyExisting, false) — Vererbung nur trennen, wenn nötig",
            "roleAssignments.add(userId, roleDefId) — Rolle setzen",
            "reset: resetRoleInheritance() — Vererbung wiederherstellen",
        ],
        "notes": [
            "getByType statt getByName ist der Lokalisierungs-Trick: Auf einer deutschen Site heißt die Rolle 'Mitwirken', nicht 'Contribute'. Der RoleType (Zahl) ist sprachunabhängig — derselbe Kniff wie in der C#-Lösung.",
            "ensureUser: Bevor man jemandem etwas zuweisen kann, muss der Benutzer in der versteckten User-Information-List der Site existieren. ensureUser legt ihn bei Bedarf an und liefert die numerische userId.",
            "breakRoleInheritance nur, wenn das Element noch erbt (HasUniqueRoleAssignments prüfen). copyExisting steuert, ob bestehende Zugriffe übernommen werden.",
            "Fehler-Mapping zeigen: PnPjs wirft HttpRequestError mit .status. 401/403 von SharePoint bilden wir auf 502 ab — das heißt 'die MI-Berechtigung fehlt', nicht 'der Aufrufer ist schuld'.",
        ],
    },

    # 18. Schritt 7 — Handler verdrahten
    {
        "title": "Schritt 7 — Handler verdrahten",
        "bullets": [
            "functions/ManagePermissions.ts — die Hello-world-Logik ersetzen",
            "1) authorizeCaller(authHeader) → bei Fehler sofort raus",
            "2) request.json() → Body lesen (ungültig → 400)",
            "3) action grant/reset → Service aufrufen",
            "Struktur identisch zu Schritt 2 — nur die Logik ist gewachsen",
        ],
        "notes": [
            "Schön zeigen: Der Handler ist dünn. Er orchestriert nur — authentifizieren, Body lesen, an den Service delegieren. Die eigentliche Arbeit liegt in auth/ und services/.",
            "Reihenfolge ist Sicherheit: ERST Token prüfen, DANN überhaupt den Body anfassen. Ein nicht autorisierter Aufrufer kommt nie bis zur SharePoint-Logik.",
            "Vergleich zur Hello-world-Version aus Schritt 2 nebeneinanderlegen: app.http-Registrierung identisch, Handler-Signatur identisch — nur der Rumpf ist gewachsen. Das nimmt die Angst vor der Komplexität.",
        ],
    },

    # 19. Schritt 8 — Lokal testen
    {
        "title": "Schritt 8 — Lokal testen",
        "bullets": [
            "local.settings.json: AzureAd__* + SharePoint__AllowedHosts eintragen",
            "az login — lokal nutzt DefaultAzureCredential diese Anmeldung",
            "azurite (Terminal 1) + npm start (Terminal 2)",
            "Token holen: az account get-access-token --resource api://<client-id>",
            "Invoke-RestMethod mit Bearer-Token gegen localhost:7071",
        ],
        "notes": [
            "Der Clou: Dank DefaultAzureCredential testet man lokal gegen ECHTES SharePoint, ohne irgendetwas zu deployen — es zählt die az-login-Anmeldung.",
            "local.settings.json wird nicht deployt; sie ist nur für die lokale Entwicklung. In Azure kommen dieselben Schlüssel als App Settings.",
            "Hinweis zum Test-Token: Das az-CLI-Token enthält access_as_user nur, wenn Tenant-Richtlinien das zulassen. Für realistische Tests später über das Web-Part gehen.",
            "Wer hier grün sieht, hat die ganze Logik bewiesen — der Rest ist Deployment und Berechtigungen.",
        ],
    },

    # 20. Schritt 9 — Deployen
    {
        "title": "Schritt 9 — Nach Azure deployen",
        "bullets": [
            "az group create + az storage account create",
            "az functionapp create --flexconsumption-location … --runtime node --runtime-version 22",
            "az functionapp identity assign — System-assigned Managed Identity an",
            "az functionapp config appsettings set — AzureAd__* + SharePoint__AllowedHosts",
            "az functionapp cors add (SPFx-Origin) · func azure functionapp publish",
        ],
        "notes": [
            "Flex Consumption Plan: aktuell empfohlen — serverless, scale-to-zero, Node 22, schnelles Scale-out. Pay-per-execution.",
            "identity assign liefert die principalId der MI zurück — die brauchen wir in Schritt 11. Mitschreiben!",
            "App Settings sind die Cloud-Variante von local.settings.json. In Node liest man sie direkt über process.env — das __ bleibt im Variablennamen erhalten.",
            "CORS nicht vergessen, sonst lehnt der Browser den Aufruf des Web-Parts ab. Und: func publish macht auf Linux/Flex einen Remote-Build (npm install + tsc in Azure).",
            "Flex-Falle erwähnen: Der echte Host hat ein Region-/Hash-Suffix — immer defaultHostName abfragen, nicht <app>.azurewebsites.net annehmen.",
        ],
    },

    # 21. Schritt 10 — Entra-App + Gruppe
    {
        "title": "Schritt 10 — Entra-App + Sicherheitsgruppe",
        "bullets": [
            "App-Registrierung 'ManagePermissions API' → Application (client) ID",
            "Expose an API → Scope access_as_user",
            "Token-Version 2 (Manifest) + groups claim (Token configuration)",
            "Sicherheitsgruppe anlegen → Object Id = AzureAd__AllowedGroupId",
            "Gruppe der Enterprise-App zuweisen (sonst kein groups-Claim)",
        ],
        "notes": [
            "Das ist die Konfiguration der ERSTEN Vertrauensgrenze — und sie ist sprachunabhängig, identisch zur C#-Lösung. Es geht um Identität und Token, nicht um Code.",
            "Hier schlagen die beiden Stolpersteine von Folie 15 zu: Token v2 und groups-Claim. Beide JETZT richtig setzen, dann später kein 401/403-Rätselraten.",
            "Die Gruppe der Enterprise-App zuweisen ist der oft vergessene zweite Schritt — ohne ihn erscheint der groups-Claim nicht, selbst wenn man ihn aktiviert hat.",
            "Client-ID und Group-Object-ID müssen als App Settings gesetzt sein (Schritt 9).",
        ],
    },

    # 22. Schritt 11 — MI berechtigen
    {
        "title": "Schritt 11 — Managed Identity berechtigen",
        "bullets": [
            "Least privilege: Sites.Selected + per-Site FullControl",
            "PnPjs/CSOM nutzt das SharePoint-Token → Sites.Selected auf Graph UND SPO nötig",
            "App-Rollen via Microsoft.Graph: New-MgServicePrincipalAppRoleAssignment",
            "Per-Site: Grant-PnPAzureADAppSitePermission … -Permissions FullControl",
            "write reicht NICHT für Item-Berechtigungen (laut MS-Doku)",
        ],
        "notes": [
            "Das betrifft die Identität der Function App, nicht den Code — exakt dieselben Schritte wie bei der C#-Lösung.",
            "Wichtige Falle: PnPjs/CSOM spricht die SharePoint-REST-API mit einem SharePoint-Token. Graph-Sites.Selected allein reicht NICHT — man braucht die App-Rolle Sites.Selected auch auf der SharePoint-Online-API.",
            "Connect-PnPOnline braucht seit PnP.PowerShell 2.x eine eigene -ClientId: Das Modul bringt KEINE vorregistrierte Anmelde-App mehr mit. Einmalig pro Tenant mit Register-PnPEntraIDAppForInteractiveLogin eine Public-Client-App (Redirect http://localhost) anlegen, die zurückgegebene ClientId danach bei -Interactive -ClientId wiederverwenden. Ohne -ClientId schlägt das interaktive Login fehl.",
            "Danach die konkrete Per-Site-Berechtigung FullControl. write klingt sparsamer, reicht aber laut Microsoft-Doku für das Verwalten von Item-Berechtigungen nicht.",
            "Geduld bei 401: Nach Rollenänderung kann der MI-Token-Cache bis ~24 h alt sein. Hilft: az functionapp stop + start, oder abwarten.",
        ],
    },

    # 23. C# vs TypeScript (Tabelle)
    {
        "title": "C# vs. TypeScript — gegenübergestellt",
        "bullets": [
            "Gleiche Architektur, andere Werkzeuge:",
        ],
        "table": {
            "columns": ["Baustein", "C# / TypeScript"],
            "rows": [
                ["Function registrieren", "[HttpTrigger]-Attribut  /  app.http(...)"],
                ["Token validieren", "Microsoft.IdentityModel  /  jose"],
                ["SharePoint-Zugriff", "PnP Core SDK  /  PnPjs (spfi)"],
                ["App-only-Token", "DefaultAzureCredential  /  AzureIdentity + DefaultAzureCredential"],
                ["Konfiguration", "IOptions<T>  /  process.env"],
                ["DI / Komposition", "Program.cs (CreateBuilder)  /  import + Modul-Singletons"],
                ["Vererbung trennen", "BreakRoleInheritanceAsync  /  breakRoleInheritance"],
                ["Deployen", "func azure functionapp publish  (beide)"],
            ],
        },
        "notes": [
            "Diese Tabelle nimmt C#-Entwicklern die Scheu und zeigt TypeScript-Leuten, dass es kein Hexenwerk ist. Jede Zeile ist eine 1:1-Entsprechung.",
            "DI / Komposition erklären, falls gefragt: DI = Dependency Injection ('Abhängigkeiten einspritzen'). Eine zentrale Stelle erzeugt die Bausteine und steckt sie zusammen (komponiert), statt dass jeder Baustein seine Helfer selbst baut. C# hat dafür einen eingebauten Mechanismus in Program.cs (FunctionsApplication.CreateBuilder). In TypeScript brauchen wir kein Framework: Wir verbinden die Teile über import und legen Zustand einmal auf Modulebene an (z.B. die eine DefaultAzureCredential). Gleiche Idee, weniger Zeremonie.",
            "Was sich NICHT ändert: App-Registrierung, der Token-v2-/Gruppen-Claim-Stolperstein und das Sites.Selected-Modell der Managed Identity. Das ist Plattform, nicht Sprache.",
            "Botschaft: Die Sprache ist austauschbar, das Muster bleibt. Wer es einmal in einer Sprache verstanden hat, überträgt es in jede andere.",
        ],
    },

    # 24. Übertragbarkeit
    {
        "title": "Übertragbarkeit — das eigentliche Lernziel",
        "bullets": [
            "Drei Fragen für JEDEN Azure-Code, der auf M365 zugreift:",
            "   1) Welche Identität führt meinen Code aus?",
            "   2) Woher kommt das Token — und was steht drin?",
            "   3) Wie greife ich least-privilege zu (Sites.Selected statt FullControl)?",
            "Gilt für C#, TypeScript, PowerShell, Python — und jede Azure-Workload",
        ],
        "notes": [
            "Das ist die Zusammenfassung, die hängenbleiben soll. Nicht die einzelnen func-Befehle, sondern diese drei Fragen.",
            "Wer diese drei Fragen für eine Function beantworten kann, kann es auch für einen Worker, eine Logic App, eine CI-Pipeline oder Code in einer anderen Cloud.",
            "Das Upload-Counter-Beispiel und ManagePermissions waren nur Vehikel — das Auth- und Permission-Gerüst ist überall dasselbe.",
            "Tauscht in Gedanken ManagePermissions gegen euren echten Use-Case aus: Das Gerüst bleibt identisch.",
        ],
    },

    # 25. Danke
    {
        "title": "Vielen Dank!",
        "bullets": [
            "Anleitung zum Mitmachen: anleitung.md (Befehl für Befehl, kompletter Code)",
            "Produktive C#-Variante: ManagePermissions/ (Schwester-Repo)",
            "Konzept-Deep-Dive: praesentation/azure-functions-konzept.md",
            "Fragen?",
        ],
        "notes": [
            "anleitung.md ist self-contained — die Teilnehmer können zuhause Schritt für Schritt nachbauen.",
            "Wer C# bevorzugt: Die produktive ManagePermissions-Lösung liegt im Schwester-Repo, gleiche Architektur.",
            "azure-functions-konzept.md für alle, die mehr zu Triggern, Bindings und Hosting-Modellen wollen.",
            "Fragen sammeln, dann Übergang zum Hands-On oder zur Pause.",
        ],
    },
]

# ---------------------------------------------------------------------------
# Theme: Azure-blau, klare Hierarchie, kein Branding/Logo (identisch zum Haupt-Deck).
# ---------------------------------------------------------------------------
COLOR_PRIMARY = RGBColor(0x00, 0x78, 0xD4)  # Azure Blue
COLOR_DARK    = RGBColor(0x1F, 0x2D, 0x54)  # Deep Navy (Header-/Title-Block)
COLOR_ACCENT  = RGBColor(0x50, 0xE6, 0xFF)  # Cyan (Workshop-Label)
COLOR_TEXT    = RGBColor(0x24, 0x24, 0x24)
COLOR_MUTED   = RGBColor(0x59, 0x59, 0x59)
COLOR_RULE    = RGBColor(0xD0, 0xD7, 0xDE)
COLOR_WHITE   = RGBColor(0xFF, 0xFF, 0xFF)

# 16:9-Layout. Alle Folien aus dem Blank-Layout selbst aufgebaut, damit die
# 4:3-Default-Placeholder von python-pptx nicht stören.
SLIDE_WIDTH       = Inches(13.333)
SLIDE_HEIGHT      = Inches(7.5)
MARGIN_X          = Inches(0.7)
CONTENT_WIDTH     = SLIDE_WIDTH - 2 * MARGIN_X
HEADER_HEIGHT     = Inches(1.1)
ACCENT_HEIGHT     = Inches(0.08)
BODY_TOP          = HEADER_HEIGHT + Inches(0.4)
FOOTER_Y          = Inches(7.0)
BODY_BOTTOM       = FOOTER_Y - Inches(0.15)

TITLE_FONT_SIZE   = Pt(34)
BULLET_FONT_SIZE  = Pt(22)
SUB_BULLET_SIZE   = Pt(19)
INTRO_TITLE_SIZE  = Pt(54)
INTRO_LABEL_SIZE  = Pt(14)
INTRO_BULLET_SIZE = Pt(24)
FOOTER_FONT_SIZE  = Pt(10)

FONT_TITLE        = "Segoe UI Semibold"
FONT_BODY         = "Segoe UI"
WORKSHOP_LABEL    = "Azure Function von Grund auf · Developer Workshop 2026"

BULLET_GLYPH      = "▸"
SUB_PREFIX_HINTS  = ("→", " ", "\t")

# Tabellen (Referenz-Folien)
COLOR_OK          = RGBColor(0x13, 0x73, 0x33)  # Häkchen-Grün
COLOR_NO          = RGBColor(0xB8, 0xC0, 0xC8)  # dezentes Grau
TABLE_ROW_ALT     = RGBColor(0xEF, 0xF4, 0xFA)  # Zebra-Streifen hell
TABLE_HEADER_SIZE = Pt(11)
TABLE_LABEL_SIZE  = Pt(10)
TABLE_CELL_SIZE   = Pt(10)
CAPTION_FONT_SIZE = Pt(13)


def _blank_slide(prs: Presentation):
    return prs.slides.add_slide(prs.slide_layouts[6])  # Blank


def _add_rect(slide, left, top, width, height, fill_color):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    shape.line.fill.background()
    try:
        shape.shadow.inherit = False
    except AttributeError:
        pass
    return shape


def _new_textbox(slide, left, top, width, height, *, vertical=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.margin_left = Emu(0)
    tf.margin_right = Emu(0)
    tf.margin_top = Emu(0)
    tf.margin_bottom = Emu(0)
    tf.word_wrap = True
    tf.vertical_anchor = vertical
    return box


def _set_run(run, text, *, font, size, color, bold=False):
    run.text = text
    run.font.name = font
    run.font.size = size
    run.font.color.rgb = color
    run.font.bold = bold


def _add_header(slide, title: str) -> None:
    _add_rect(slide, Emu(0), Emu(0), SLIDE_WIDTH, HEADER_HEIGHT, COLOR_DARK)
    _add_rect(slide, Emu(0), HEADER_HEIGHT, SLIDE_WIDTH, ACCENT_HEIGHT, COLOR_PRIMARY)
    box = _new_textbox(
        slide, MARGIN_X, Emu(0),
        CONTENT_WIDTH, HEADER_HEIGHT,
        vertical=MSO_ANCHOR.MIDDLE,
    )
    p = box.text_frame.paragraphs[0]
    _set_run(p.add_run(), title, font=FONT_TITLE, size=TITLE_FONT_SIZE,
             color=COLOR_WHITE, bold=True)


def _add_footer(slide, page: int, total: int) -> None:
    _add_rect(slide, MARGIN_X, FOOTER_Y, CONTENT_WIDTH, Emu(9525), COLOR_RULE)
    label_box = _new_textbox(slide, MARGIN_X, FOOTER_Y + Inches(0.08),
                             Inches(9), Inches(0.3))
    _set_run(label_box.text_frame.paragraphs[0].add_run(),
             WORKSHOP_LABEL, font=FONT_BODY, size=FOOTER_FONT_SIZE, color=COLOR_MUTED)
    page_box = _new_textbox(slide, SLIDE_WIDTH - MARGIN_X - Inches(2),
                            FOOTER_Y + Inches(0.08), Inches(2), Inches(0.3))
    p = page_box.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    _set_run(p.add_run(), f"{page} / {total}", font=FONT_BODY,
             size=FOOTER_FONT_SIZE, color=COLOR_MUTED)


def _is_sub_bullet(line: str) -> bool:
    return line.startswith("   ") or line.lstrip().startswith(SUB_PREFIX_HINTS[0])


def _add_body(slide, bullets: list) -> None:
    body_height = BODY_BOTTOM - BODY_TOP
    box = _new_textbox(slide, MARGIN_X, BODY_TOP, CONTENT_WIDTH, body_height)
    tf = box.text_frame
    for i, line in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if i > 0:
            p.space_before = Pt(8)
        if _is_sub_bullet(line):
            _set_run(p.add_run(), "     ", font=FONT_BODY, size=SUB_BULLET_SIZE,
                     color=COLOR_MUTED)
            _set_run(p.add_run(), line.lstrip(), font=FONT_BODY,
                     size=SUB_BULLET_SIZE, color=COLOR_MUTED)
        else:
            _set_run(p.add_run(), f"{BULLET_GLYPH}  ", font=FONT_BODY,
                     size=BULLET_FONT_SIZE, color=COLOR_PRIMARY, bold=True)
            _set_run(p.add_run(), line, font=FONT_BODY,
                     size=BULLET_FONT_SIZE, color=COLOR_TEXT)


def _set_speaker_notes(slide, notes: list) -> None:
    if not notes:
        return
    # Inline-Aufzählungen (" 1)", " 2)", ...) auf eigene Zeile umbrechen.
    rendered = [re.sub(r" (\d+\)) ", r"\n\1 ", line) for line in notes]
    slide.notes_slide.notes_text_frame.text = "\n".join(rendered)


def _add_caption(slide, lines: list):
    """Kurze Erläuterungs-/Legenden-Zeilen über einer Tabelle. Gibt die
    y-Position zurück, ab der die Tabelle beginnen darf."""
    line_h = Inches(0.32)
    box = _new_textbox(slide, MARGIN_X, BODY_TOP, CONTENT_WIDTH, line_h * len(lines))
    tf = box.text_frame
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        _set_run(p.add_run(), line, font=FONT_BODY, size=CAPTION_FONT_SIZE, color=COLOR_MUTED)
    return BODY_TOP + line_h * len(lines) + Inches(0.12)


def _style_cell(cell, fill):
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    cell.margin_left = Inches(0.06)
    cell.margin_right = Inches(0.04)
    cell.margin_top = Emu(0)
    cell.margin_bottom = Emu(0)
    cell.fill.solid()
    cell.fill.fore_color.rgb = fill


def _add_table(slide, columns: list, rows: list, top) -> None:
    n_rows = len(rows) + 1
    n_cols = len(columns)
    height = BODY_BOTTOM - top
    gfx = slide.shapes.add_table(n_rows, n_cols, MARGIN_X, top, CONTENT_WIDTH, height)
    table = gfx.table
    table.first_row = False      # eigenes Header-Styling, kein Theme-Banding
    table.horz_banding = False

    first_w = Inches(4.3)
    rest_w = Emu(int((CONTENT_WIDTH - first_w) / (n_cols - 1)))
    table.columns[0].width = first_w
    for c in range(1, n_cols):
        table.columns[c].width = rest_w

    row_h = Emu(int(height / n_rows))
    for r in range(n_rows):
        table.rows[r].height = row_h

    # Header-Zeile
    for c, label in enumerate(columns):
        cell = table.cell(0, c)
        _style_cell(cell, COLOR_DARK)
        cell.text_frame.word_wrap = True
        p = cell.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT if c == 0 else PP_ALIGN.CENTER
        _set_run(p.add_run(), label, font=FONT_TITLE, size=TABLE_HEADER_SIZE,
                 color=COLOR_WHITE, bold=True)

    # Datenzeilen
    for r, row in enumerate(rows, start=1):
        fill = TABLE_ROW_ALT if (r % 2 == 0) else COLOR_WHITE
        for c, val in enumerate(row):
            cell = table.cell(r, c)
            _style_cell(cell, fill)
            p = cell.text_frame.paragraphs[0]
            if c == 0:
                p.alignment = PP_ALIGN.LEFT
                _set_run(p.add_run(), val, font=FONT_BODY, size=TABLE_LABEL_SIZE,
                         color=COLOR_TEXT)
            else:
                p.alignment = PP_ALIGN.LEFT
                _set_run(p.add_run(), val, font=FONT_BODY, size=TABLE_CELL_SIZE,
                         color=COLOR_TEXT)


def add_title_slide(prs: Presentation, title: str, subtitle_lines: list,
                    notes: list | None = None) -> None:
    slide = _blank_slide(prs)
    block_width = Inches(5.8)
    _add_rect(slide, Emu(0), Emu(0), block_width, SLIDE_HEIGHT, COLOR_DARK)
    _add_rect(slide, block_width, Emu(0), Inches(0.15), SLIDE_HEIGHT, COLOR_PRIMARY)

    label_box = _new_textbox(slide, Inches(0.7), Inches(0.8),
                             block_width - Inches(1.4), Inches(0.4))
    _set_run(label_box.text_frame.paragraphs[0].add_run(),
             "WORKSHOP", font=FONT_TITLE, size=INTRO_LABEL_SIZE,
             color=COLOR_ACCENT, bold=True)

    title_box = _new_textbox(slide, Inches(0.7), Inches(2.4),
                             block_width - Inches(1.4), Inches(3.5))
    p = title_box.text_frame.paragraphs[0]
    _set_run(p.add_run(), title, font=FONT_TITLE,
             size=INTRO_TITLE_SIZE, color=COLOR_WHITE, bold=True)

    sub_left = block_width + Inches(0.6)
    sub_box = _new_textbox(slide, sub_left, Inches(2.4),
                           SLIDE_WIDTH - sub_left - Inches(0.7), Inches(4.2))
    for i, line in enumerate(subtitle_lines):
        p = sub_box.text_frame.paragraphs[0] if i == 0 else sub_box.text_frame.add_paragraph()
        if i > 0:
            p.space_before = Pt(14)
        _set_run(p.add_run(), line, font=FONT_BODY,
                 size=INTRO_BULLET_SIZE, color=COLOR_TEXT)

    foot_box = _new_textbox(slide, sub_left, Inches(6.8),
                            SLIDE_WIDTH - sub_left - Inches(0.7), Inches(0.3))
    _set_run(foot_box.text_frame.paragraphs[0].add_run(),
             WORKSHOP_LABEL, font=FONT_BODY, size=FOOTER_FONT_SIZE, color=COLOR_MUTED)
    _set_speaker_notes(slide, notes)


def add_content_slide(prs: Presentation, title: str, bullets: list,
                     page: int, total: int, notes: list | None = None,
                     table: dict | None = None) -> None:
    slide = _blank_slide(prs)
    _add_header(slide, title)
    if table:
        top = _add_caption(slide, bullets) if bullets else BODY_TOP
        _add_table(slide, table["columns"], table["rows"], top)
    else:
        _add_body(slide, bullets)
    _add_footer(slide, page, total)
    _set_speaker_notes(slide, notes)


def main() -> None:
    parser = argparse.ArgumentParser(description="Build the 'Azure Function from scratch' slide deck.")
    parser.add_argument(
        "--output", "-o",
        default="slides.pptx",
        help="Output filename (default: slides.pptx). Use a different name if the file is locked.",
    )
    args = parser.parse_args()

    out_path = Path(__file__).parent / args.output
    prs = Presentation()
    prs.slide_width  = SLIDE_WIDTH
    prs.slide_height = SLIDE_HEIGHT

    total = len(SLIDES)
    first = SLIDES[0]
    add_title_slide(prs, first["title"], first["bullets"], first.get("notes"))

    for idx, slide in enumerate(SLIDES[1:], start=2):
        add_content_slide(prs, slide["title"], slide["bullets"],
                          idx, total, slide.get("notes"), slide.get("table"))

    prs.save(out_path)
    print(f"Slides geschrieben: {out_path}  ({len(SLIDES)} Folien)")


if __name__ == "__main__":
    main()
